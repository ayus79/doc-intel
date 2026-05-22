import io
import struct
import xml.etree.ElementTree as ET
import pymupdf
import pdfplumber
import pytesseract
from PIL import Image
from dataclasses import dataclass

import docx
import pptx
import pandas as pd

from core.constants import (
    SUPPORTED_FORMATS,
    RASTER_IMAGE_FORMATS,
    IMAGE_UPSCALE_MIN_WIDTH,
    OCR_DPI,
    DOCX_LINES_PER_PAGE,
    TEXT_LINES_PER_PAGE,
)


@dataclass
class ParsedPage:
    doc_name: str
    page_number: int
    text: str


class DocumentParser:

    def parse(self, file_bytes: bytes, filename: str) -> list[ParsedPage]:
        ext = filename.rsplit(".", 1)[-1].lower()
        if ext not in SUPPORTED_FORMATS:
            raise ValueError(
                f"Unsupported format: .{ext}. Supported: {SUPPORTED_FORMATS}"
            )
        if ext in ("xlsx",):
            return self._parse_spreadsheet(file_bytes, filename, ext)
        if ext in RASTER_IMAGE_FORMATS:
            return self._parse_raster_image(file_bytes, filename)
        if ext == "svg":
            return self._parse_svg(file_bytes, filename)
        method = getattr(self, f"_parse_{ext}", None)
        return method(file_bytes, filename)

    # ---------- PDF ----------

    def _parse_pdf(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        pages = []
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        with pdfplumber.open(io.BytesIO(file_bytes)) as plumber_doc:
            for page_num in range(len(doc)):
                page = doc[page_num]
                table_text = self._extract_pdf_tables(plumber_doc.pages[page_num])
                plain_text = page.get_text().strip()
                text = "\n".join(filter(None, [table_text, plain_text])).strip()
                if not text:
                    text = self._ocr_page(page)
                if text:
                    pages.append(
                        ParsedPage(
                            doc_name=doc_name, page_number=page_num + 1, text=text
                        )
                    )
        doc.close()
        return pages

    def _ocr_page(self, page) -> str:
        pix = page.get_pixmap(dpi=OCR_DPI)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img).strip()

    def _extract_pdf_tables(self, plumber_page) -> str:
        lines = []
        for table in plumber_page.extract_tables():
            for row in table:
                clean = [cell.strip() if cell else "" for cell in row]
                lines.append(" | ".join(clean))
        return "\n".join(lines)

    # ---------- DOCX ----------

    def _parse_docx(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        document = docx.Document(io.BytesIO(file_bytes))
        pages, current, page_num = [], [], 1

        for para in document.paragraphs:
            text = para.text.strip()
            if text:
                current.append(text)
            if len(current) >= DOCX_LINES_PER_PAGE:
                pages.append(
                    ParsedPage(
                        doc_name=doc_name, page_number=page_num, text="\n".join(current)
                    )
                )
                current, page_num = [], page_num + 1
        if current:
            pages.append(
                ParsedPage(
                    doc_name=doc_name, page_number=page_num, text="\n".join(current)
                )
            )

        for i, table in enumerate(document.tables):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            if rows:
                pages.append(
                    ParsedPage(
                        doc_name=doc_name, page_number=1000 + i, text="\n".join(rows)
                    )
                )

        return pages

    # ---------- PPTX ----------

    def _parse_pptx(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        prs = pptx.Presentation(io.BytesIO(file_bytes))
        pages = []
        for i, slide in enumerate(prs.slides):
            parts = [
                para.text.strip()
                for shape in slide.shapes
                if shape.has_text_frame
                for para in shape.text_frame.paragraphs
                if para.text.strip()
            ]
            if parts:
                pages.append(
                    ParsedPage(
                        doc_name=doc_name, page_number=i + 1, text="\n".join(parts)
                    )
                )
        return pages

    # ---------- XLSX / CSV ----------

    def _parse_csv(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        return self._parse_spreadsheet(file_bytes, doc_name, "csv")

    def _parse_spreadsheet(
        self, file_bytes: bytes, doc_name: str, ext: str
    ) -> list[ParsedPage]:
        if ext == "csv":
            df_map = {"Sheet1": pd.read_csv(io.BytesIO(file_bytes))}
        else:
            df_map = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)

        pages = []
        for sheet_name, df in df_map.items():
            df = df.fillna("").astype(str)
            header = " | ".join(df.columns.tolist())
            rows = [" | ".join(row) for _, row in df.iterrows()]
            text = f"[Sheet: {sheet_name}]\n{header}\n" + "\n".join(rows)
            pages.append(ParsedPage(doc_name=doc_name, page_number=1, text=text))
        return pages

    # ---------- TXT / MD ----------

    def _parse_txt(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        return self._parse_plain_text(file_bytes, doc_name)

    def _parse_md(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        return self._parse_plain_text(file_bytes, doc_name)

    def _parse_plain_text(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        lines = file_bytes.decode("utf-8", errors="ignore").strip().splitlines()
        pages, chunk_size = [], TEXT_LINES_PER_PAGE
        for i in range(0, len(lines), chunk_size):
            chunk = "\n".join(lines[i : i + chunk_size]).strip()
            if chunk:
                pages.append(
                    ParsedPage(
                        doc_name=doc_name, page_number=i // chunk_size + 1, text=chunk
                    )
                )
        return pages

    @staticmethod
    def _strip_png_chunk(data: bytes, chunk_name: str) -> bytes:
        out = bytearray(data[:8])
        pos = 8
        while pos < len(data):
            length = struct.unpack(">I", data[pos : pos + 4])[0]
            ctype = data[pos + 4 : pos + 8].decode("ascii", errors="replace")
            if ctype != chunk_name:
                out += data[pos : pos + 12 + length]
            pos += 12 + length
        return bytes(out)

    # ---------- Raster Images (JPG, PNG, WEBP, BMP, TIFF) ----------

    def _parse_raster_image(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        if file_bytes[:8] == b"\x89PNG\r\n\x1a\n":
            file_bytes = self._strip_png_chunk(file_bytes, "eXIf")

        img = Image.open(io.BytesIO(file_bytes))
        img.load()
        img = img.convert("RGB")
        w, h = img.size
        if w < IMAGE_UPSCALE_MIN_WIDTH:
            scale = IMAGE_UPSCALE_MIN_WIDTH / w
            img = img.resize((int(w * scale), int(h * scale)), Image.LANCZOS)
        text = pytesseract.image_to_string(img).strip()
        if not text:
            return []
        return [ParsedPage(doc_name=doc_name, page_number=1, text=text)]

    # ---------- SVG ----------

    def _parse_svg(self, file_bytes: bytes, doc_name: str) -> list[ParsedPage]:
        root = ET.fromstring(file_bytes.decode("utf-8", errors="ignore"))
        ns = {"svg": "http://www.w3.org/2000/svg"}
        parts = []

        for el in root.iter():
            t = (el.text or "").strip()
            tail = (el.tail or "").strip()
            if t:
                parts.append(t)
            if tail:
                parts.append(tail)
        text = "\n".join(parts).strip()
        if not text:
            return []
        return [ParsedPage(doc_name=doc_name, page_number=1, text=text)]


parser = DocumentParser()
